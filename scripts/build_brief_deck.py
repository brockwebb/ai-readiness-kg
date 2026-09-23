#!/usr/bin/env python3
"""The brief deck, rendered from its content file and the brief's material pack.

`cc_tasks/2026-09-22_brief_deck_assembly.md`, under DN-005 (the brief is a VIEW of the
framework) and DN-008 ruling 2 (chapter order). **Zero model calls, no network.** Reads only
`docs/brief/` (the pack), `docs/deck/brief_deck_content.md` (the authored slides),
`corpus/manifest.json` (one summary slide) and git; writes only under `docs/deck/`.

    scripts/build_brief_deck.py                    render docs/deck/brief_deck.pptx and
                                                   docs/deck/brief_appendix.pptx
    scripts/build_brief_deck.py --check            re-render both to a temp path; exit 1 on drift
    scripts/build_brief_deck.py --render-diagrams  re-draw the four Mermaid diagrams with mmdc

**The renderer is `scripts/build_framework_deck.py`.** Its body parser, its layout rule (18pt
stepped to a 14pt floor, then a split on line boundaries) and its inline-bold runs are imported,
not copied. What this file adds is what the brief needs and the framework deck did not:

* **Two outputs from one content file** (`cc_tasks/2026-09-23_brief_deck_packaging.md`,
  decision 1, for DN-008 ruling 2's roughly-100-slide shape). The brief, `brief_deck.pptx`, is
  every authored section outside the `Appendix` chapter: the cover, chapters A to H and a closing
  slide whose `@stamp appendix` names the appendix file and its slide count. The appendix,
  `brief_appendix.pptx`, is its own cover, the `Appendix` chapter's authored sections and the
  generated slides. No slide is dropped by the split.
* **Two kinds of slide** (decision 2). *Authored* slides are the `## Slide N — Title` sections of
  the content file. *Generated* slides, the appendix, are rendered here from the pack at build
  time: one per indicator sheet, one per rule group, one corpus summary.
* **Directives** copy pack content onto a slide by script, so a number on a table, a CSV row or a
  captured command output is never retyped: `@table <page> | <heading> [| cols=..] [| nth=N]`,
  `@csv <file> | cols=.. [| where=col=val]`, `@capture <step id>`, `@diagram <a|b|c|d>`,
  `@provenance`, `@stamp` (date, pack commit and cycle, read from git and the pack),
  `@stamp appendix` (the same, prefixed by the appendix file and its rendered slide count).
* **Quotations are verified.** A body line `> text` must occur in one of the files the slide's
  `source:` line names, under `kg/extraction/grounding.py` normalization (the grounding test
  every edge in this repository passes). A quote that is not found refuses the build.
* **The numeral gate** (decision 4). Every numeral in a slide's title and in its lines that are
  neither a verified quote nor directive output must be a `value` in `docs/brief/numbers.json`
  under a page the slide's `source:` line names. The numeral scan is the pack's own
  (`tests/test_brief_pack.py::prose_numerals`), imported so the two gates cannot drift.
* **Byte-stable output.** python-pptx stamps each zip member with the wall-clock time, which is
  the only thing that differs between two renders of the same content (measured 2026-09-22 on the
  framework deck: every member identical, header bytes differ). The archive is rewritten with a
  fixed member timestamp, the reproducible-builds.org prescription for archives, so `--check`
  can compare bytes.

**Diagrams.** `mmdc` output is not guaranteed byte-stable across machines, so the PNGs are drawn
once by `--render-diagrams` into `docs/deck/diagrams/` with a sidecar recording the sha256 of the
Mermaid source each was drawn from. A render refuses when the pack's Mermaid no longer matches the
sidecar: a stale diagram stops the build rather than shipping.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "scripts"))

import build_framework_deck as FD  # noqa: E402
from kg.extraction.grounding import normalize  # noqa: E402

PACK = REPO / "docs" / "brief"
OUT_DIR = REPO / "docs" / "deck"
CONTENT = OUT_DIR / "brief_deck_content.md"
DECK = OUT_DIR / "brief_deck.pptx"
APPENDIX = OUT_DIR / "brief_appendix.pptx"
#: The chapter whose authored sections open the appendix file rather than the brief.
APPENDIX_CHAPTER = "Appendix"
DIAGRAMS = OUT_DIR / "diagrams"
DIAGRAM_SIDECAR = DIAGRAMS / "diagrams.json"
MANIFEST = REPO / "corpus" / "manifest.json"
LEDGER = PACK / "numbers.json"
CAPTURE = PACK / "D_demo_capture.json"
TASK = "cc_tasks/2026-09-22_brief_deck_assembly.md"

#: Zip member timestamp for a reproducible archive: the earliest a zip header can encode.
ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)
#: Rule groups replace one slide per rule when one per rule would exceed this (decision 2).
RULE_SLIDE_LIMIT = 40
#: Captured output is preformatted and never wraps in the layout model, so it is hard-wrapped
#: here at a width a 12pt monospace line holds inside the 12.33-inch body (12.33 * 72 / 7.2
#: = 123 characters at 0.6 em per character; 110 leaves margin). Wrapping moves characters to a
#: continuation line; none is dropped.
PRE_WRAP = 110
#: Diagram slides, in inches. *Stacked*: the image across the body width, the text box under
#: it. *Side*: the image on the left at full body height, the text in a column to its right.
#: A diagram takes whichever placement draws it larger, provided its text fits there
#: (2026-09-23 packaging task, decision 3: drawn upright, (c) and (d) are narrower than tall
#: enough that the stacked box would draw them no larger, or smaller, than left to right did).
DIAGRAM_IMG_H = 3.7
DIAGRAM_TEXT_H = 1.65
DIAGRAM_SIDE_H = 5.55
DIAGRAM_SIDE_TEXT_W = 5.0
DIAGRAM_GAP = 0.3
FOOTER_PT = 10
#: Four-space indent (preformatted, to `FD.parse_body`) plus a no-break space, which
#: `str.lstrip(" ")` does not remove and which renders as a space.
PRE_GUARD = "    \u00a0"
DIRECTIVE_RE = re.compile(r"^@(\w[\w-]*)\s*(.*)$")
META_RE = re.compile(r"^(source|chapter|layout):\s*(.*)$")
QUOTE_RE = re.compile(r"^(\s*)> ?(.*)$")


def _prose_numerals():
    """The pack's numeral scan, imported from its test so both gates read one definition."""
    spec = importlib.util.spec_from_file_location("_brief_pack_tests",
                                                  REPO / "tests" / "test_brief_pack.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.prose_numerals


def git_out(*args: str) -> str:
    r = subprocess.run(["git", *args], cwd=REPO, capture_output=True, text=True)
    if r.returncode != 0:
        raise SystemExit(f"FATAL: git {' '.join(args)} failed: {r.stderr.strip()}")
    return r.stdout.strip()


class DeckError(SystemExit):
    """A refusal: the build stops and names the slide and the reason."""


# ------------------------------------------------------------------------------ pack reads

def pack_path(name: str) -> Path:
    """A `source:` entry: a pack-relative name, else a repository-relative path."""
    p = PACK / name
    return p if p.is_file() else REPO / name


def md_tables(text: str) -> list:
    """`[(heading, [rows])]` for every markdown table, each row a list of cells, with the
    nearest heading above it. Escaped pipes are restored."""
    out, heading, cur = [], "", None
    for line in text.splitlines():
        if line.startswith("#"):
            heading = line.lstrip("#").strip()
        if line.startswith("|"):
            if cur is None:
                cur = (heading, [])
                out.append(cur)
            if re.fullmatch(r"\|(-+\|)+", line.strip()):
                continue
            cells = re.split(r"(?<!\\)\|", line.strip())[1:-1]
            cur[1].append([c.replace("\\|", "|").strip() for c in cells])
        else:
            cur = None
    return out


def _args(arg: str) -> tuple[list, dict]:
    parts = [p.strip() for p in arg.split(" | ")]
    pos = [p for p in parts if "=" not in p.split(" ")[0]]
    kw = dict(p.split("=", 1) for p in parts if "=" in p.split(" ")[0])
    return pos, kw


def rows_as_lines(head: list, rows: list, cols: list | None) -> list:
    """Table rows as slide bullets: `**header**` once, then one bullet per row with the chosen
    cells joined by ` · `. Empty cells are left out of the row, never replaced."""
    idx = [head.index(c) for c in cols] if cols else list(range(len(head)))
    out = ["**" + " · ".join(head[i] for i in idx) + "**"]
    for r in rows:
        cells = [r[i] for i in idx if i < len(r) and r[i]]
        out.append("- " + " · ".join(cells))
    return out


def directive(name: str, arg: str, n: int, appendix_slides: int | None = None
              ) -> tuple[list, str | None]:
    """`(lines, diagram key or None)`. Lines are markdown for `FD.parse_body`.
    `appendix_slides` is the appendix file's rendered slide count, for `@stamp appendix`."""
    pos, kw = _args(arg)
    cols = kw["cols"].split(",") if "cols" in kw else None
    if name == "table":
        page, heading = pos[0], pos[1]
        tabs = [t for t in md_tables(pack_path(page).read_text(encoding="utf-8"))
                if heading in t[0]]
        nth = int(kw.get("nth", "1"))
        if len(tabs) < nth:
            raise DeckError(f"slide {n}: no table {nth} under a heading containing "
                            f"{heading!r} in {page}")
        rows = tabs[nth - 1][1]
        if cols and not set(cols) <= set(rows[0]):
            raise DeckError(f"slide {n}: {page} table has no column(s) "
                            f"{sorted(set(cols) - set(rows[0]))}")
        return rows_as_lines(rows[0], rows[1:], cols), None
    if name == "csv":
        text = pack_path(pos[0]).read_text(encoding="utf-8")
        rows = list(csv.reader(l for l in text.splitlines() if not l.startswith("#")))
        head, body = rows[0], rows[1:]
        if "where" in kw:
            col, val = kw["where"].split("=", 1)
            body = [r for r in body if r[head.index(col)] == val]
        if not body:
            raise DeckError(f"slide {n}: {pos[0]} has no row where {kw.get('where')}")
        return rows_as_lines(head, body, cols), None
    if name == "capture":
        cap = json.loads(CAPTURE.read_text(encoding="utf-8"))
        step = next((s for s in cap["steps"] if s["id"] == pos[0]), None)
        if step is None:
            raise DeckError(f"slide {n}: no captured step {pos[0]!r} in {CAPTURE.name}")
        out = [f"Output: exit {step['exit']}, {step['lines_total']} lines, "
               f"run at {step['started_at']}."]
        shown = step["head"] + (["[…]"] + step["tail"] if step["tail"] else [])
        # PRE_GUARD keeps a captured line that begins `- ` or `1. ` from being read as a list
        # item by `FD.parse_body`, which tests for a bullet marker before it tests the indent.
        out += [PRE_GUARD + "$ " + l for l in _hard_wrap(step["run"])]
        out += [PRE_GUARD + l for raw in shown for l in _hard_wrap(raw)]
        return out, None
    if name == "diagram":
        return [], pos[0]
    if name == "provenance":
        return provenance_lines(), None
    if name == "stamp":
        (sha_, date), h = pack_commit(), pack_header()
        stamp = f"{date} · pack commit {sha_[:12]} · cycle of record {h['cycle']}"
        pos = [x for x in pos if x]
        if pos == ["appendix"]:
            if appendix_slides is None:
                raise DeckError(f"slide {n}: @stamp appendix outside a two-output render")
            return [f"- {APPENDIX.relative_to(REPO)}: {appendix_slides} slides · {stamp}"], None
        if pos:
            raise DeckError(f"slide {n}: @stamp takes no argument but `appendix`")
        return [stamp], None
    raise DeckError(f"slide {n}: unknown directive @{name}")


def _hard_wrap(line: str) -> list:
    line = line.rstrip() or " "
    return [line[i:i + PRE_WRAP] for i in range(0, len(line), PRE_WRAP)]


def pack_header() -> dict:
    """The pack's own generator comment (INDEX.md line 1), parsed: record commit and cycle."""
    first = (PACK / "INDEX.md").read_text(encoding="utf-8").splitlines()[0]
    m = re.search(r"framework record at commit (\w+) \(([^,]+), generated_from ([^)]+)\); "
                  r"cycle of record (\S+?)\.", first)
    if not m:
        raise DeckError(f"FATAL: {PACK / 'INDEX.md'} line 1 is not the pack's generator comment")
    return {"record_commit": m.group(1), "record": m.group(2), "generated_from": m.group(3),
            "cycle": m.group(4)}


def pack_commit() -> tuple[str, str]:
    """`(sha, date)` of the last commit that wrote `docs/brief/`: the pack the deck reads."""
    out = git_out("log", "-1", "--format=%H %cs", "--", "docs/brief")
    sha, date = out.split()
    return sha, date


def provenance_lines() -> list:
    h, (sha, date) = pack_header(), pack_commit()
    return [f"- Pack: `docs/brief/` at commit `{sha[:12]}` ({date}), written by "
            "`scripts/build_brief_pack.py` and by nothing else",
            f"- Cycle of record: `{h['cycle']}`",
            f"- Framework record: `{h['record']}` at commit `{h['record_commit']}`, generated "
            f"from `{h['generated_from']}`",
            f"- Deck: `docs/deck/brief_deck_content.md` rendered by "
            "`scripts/build_brief_deck.py`; the appendix is generated from the pack at build time"]


# ------------------------------------------------------------------------------ authored

def authored(src: str, appendix_slides: int | None = None) -> list:
    """Parse the content file into slide specs, expanding directives, verifying quotes and
    running the numeral gate. Refuses on the first defect with the slide number."""
    prose_numerals = _prose_numerals()
    ledger = json.loads(LEDGER.read_text(encoding="utf-8"))
    specs, chapter = [], None
    for raw in FD.parse(src):
        n = raw["n"]
        meta, keep = {}, []
        for line in raw["body"].strip("\n").splitlines():
            m = META_RE.match(line.strip())
            if m:
                meta[m.group(1)] = m.group(2).strip()
            else:
                keep.append(line)
        sources = [s.strip() for s in meta.get("source", "").split(",") if s.strip()]
        for s in sources:
            if not pack_path(s).is_file():
                raise DeckError(f"slide {n}: source {s!r} does not exist")
        if "chapter" in meta:
            chapter = meta["chapter"]
        texts = {s: normalize(pack_path(s).read_text(encoding="utf-8")) for s in sources}
        body, checked, diagram = [], [raw["title"]], None
        for line in keep:
            dm = DIRECTIVE_RE.match(line.strip())
            if dm:
                lines, dkey = directive(dm.group(1), dm.group(2), n, appendix_slides)
                body += lines
                diagram = dkey or diagram
                continue
            qm = QUOTE_RE.match(line)
            if qm:
                inner = qm.group(2)
                probe = re.sub(r"^([-*]|\d+\.)\s+", "", inner.strip())
                if probe and not any(normalize(probe) in t for t in texts.values()):
                    raise DeckError(f"slide {n}: quote not found in {sources or 'no source'}: "
                                    f"{probe[:90]!r}")
                body.append(qm.group(1) + inner)
                continue
            body.append(line)
            checked.append(line)
        held = {x["value"] for s in sources for x in ledger.get(s, [])}
        for text in checked:
            # A backticked span with no letter in it would hide a number from the scan; it is
            # unwrapped first, so only identifiers (`CC BY 4.0`, `RULE-A1-v4`) stay exempt.
            text = re.sub(r"`([^`A-Za-z]*)`", r"\1", text)
            for num in prose_numerals(text.lstrip("#>| ")):
                if num not in held:
                    raise DeckError(f"slide {n}: numeral {num!r} is not on numbers.json under "
                                    f"{sources or 'no source'}; copy it by directive or quote, "
                                    "or file the gap against the pack")
        needs_source = any(prose_numerals(re.sub(r"`([^`A-Za-z]*)`", r"\1", t).lstrip("#>| "))
                           for t in checked)
        if needs_source and not sources:
            raise DeckError(f"slide {n}: states a number and names no source")
        specs.append({"n": n, "title": raw["title"], "body": "\n".join(body),
                      "sources": sources, "chapter": chapter,
                      "layout": meta.get("layout", "text"), "diagram": diagram,
                      "kind": "authored"})
    return specs


# ------------------------------------------------------------------------------ generated

def md_to_body(text: str) -> tuple[str, str]:
    """A pack page as `(title, slide body)`. Headings become bold lines, tables become one
    bullet per row, block quotes and prose are kept; the generator comment is dropped."""
    title, out, tab_head = "", [], None
    for line in text.splitlines():
        if line.startswith("<!--"):
            continue
        if line.startswith("# "):
            title = line[2:].strip()
            continue
        if line.startswith("#"):
            out.append("**" + line.lstrip("#").strip() + "**")
            tab_head = None
            continue
        if line.startswith("|"):
            if re.fullmatch(r"\|(-+\|)+", line.strip()):
                continue
            cells = [c.replace("\\|", "|").strip()
                     for c in re.split(r"(?<!\\)\|", line.strip())[1:-1]]
            if tab_head is None:
                tab_head = cells
                if not (len(cells) == 2 and cells[0] in ("field", "box")):
                    out.append("**" + " · ".join(cells) + "**")
                continue
            if len(tab_head) == 2 and tab_head[0] in ("field", "box"):
                out.append(f"- {cells[0]}: {cells[1]}")
            else:
                out.append("- " + " · ".join(c for c in cells if c))
            continue
        tab_head = None
        if line.startswith(">"):
            line = line.lstrip("> ").rstrip()
            if not line:
                continue
        out.append(line)
    return title, "\n".join(out)


def rule_groups() -> list:
    """`[(leg, [rows])]` from `appendix/rules.md`, in registry order."""
    tabs = md_tables((PACK / "appendix" / "rules.md").read_text(encoding="utf-8"))
    head, rows = tabs[0][1][0], tabs[0][1][1:]
    groups: dict = {}
    for r in rows:
        groups.setdefault(r[head.index("leg")], []).append(r)
    return head, list(groups.items())


def corpus_summary() -> str:
    """One slide: admitted and screened-out counts by `identity.doc_type`. The manifest has no
    `source_type` field; `doc_type` is the field page C prints, for the same reason."""
    entries = json.loads(MANIFEST.read_text(encoding="utf-8"))["entries"].values()
    by = Counter((e["identity"].get("doc_type") or "none",
                  (e.get("screening") or {}).get("decision") or "none") for e in entries)
    decisions = sorted({d for _, d in by})
    lines = ["**doc_type · " + " · ".join(decisions) + "**"]
    for t in sorted({t for t, _ in by}):
        lines.append(f"- {t} · " + " · ".join(f"{d} {by[(t, d)]}" for d in decisions))
    lines.append(f"- all · " + " · ".join(
        f"{d} {sum(v for (t, dd), v in by.items() if dd == d)}" for d in decisions))
    lines += ["", "From `corpus/manifest.json` (`identity.doc_type` by `screening.decision`). "
              "The admitted documents, one row per indicator citation, are in "
              "`docs/brief/C_provenance.csv`."]
    return "\n".join(lines)


def generated(start: int) -> list:
    specs, n = [], start
    sheets = sorted((PACK / "appendix").glob("indicator_*.md"),
                    key=lambda p: FD_code_key(p.stem[len("indicator_"):]))
    for p in sheets:
        # The sheet's own heading is the indicator text, which can run to several lines; it
        # opens the body in bold and the slide title carries the code alone.
        title, body = md_to_body(p.read_text(encoding="utf-8"))
        code = p.stem[len("indicator_"):]
        body = f"**{title}**\n\n{body}"
        specs.append({"n": n, "title": f"Appendix · Indicator {code}", "body": body,
                      "sources": [str(p.relative_to(PACK))], "chapter": "Appendix",
                      "layout": "text", "diagram": None, "kind": "indicator"})
        n += 1
    head, groups = rule_groups()
    per_rule = sum(len(r) for _, r in groups)
    cols = ["rule_id", "state", "candidate", "claim", "measures", "summary"]
    if per_rule <= RULE_SLIDE_LIMIT:
        units = [(r[head.index("rule_id")], [r]) for _, rows in groups for r in rows]
    else:
        units = groups
    for key, rows in units:
        body = "\n".join(rows_as_lines(head, rows, cols))
        specs.append({"n": n, "title": f"Appendix · Rules: {key}", "body": body,
                      "sources": ["appendix/rules.md"], "chapter": "Appendix",
                      "layout": "text", "diagram": None, "kind": "rule"})
        n += 1
    specs.append({"n": n, "title": "Appendix · The corpus, by document type",
                  "body": corpus_summary(), "sources": ["corpus/manifest.json"],
                  "chapter": "Appendix", "layout": "text", "diagram": None, "kind": "corpus"})
    return specs


def FD_code_key(code: str) -> tuple:
    m = re.match(r"([A-Z])(\d+)(.*)", code)
    return (m.group(1), int(m.group(2)), m.group(3)) if m else (code, 0, "")


# ------------------------------------------------------------------------------ diagrams

def mermaid_blocks() -> dict:
    text = (PACK / "E_architecture.md").read_text(encoding="utf-8")
    heads = re.findall(r"^## \((\w)\)", text, re.M)
    blocks = re.findall(r"```mermaid\n(.*?)```", text, re.S)
    if len(heads) != len(blocks):
        raise DeckError("FATAL: E_architecture.md headings and Mermaid blocks do not pair up")
    return dict(zip(heads, blocks))


def sha(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def render_diagrams() -> int:
    mmdc = shutil.which("mmdc")
    if mmdc is None:
        raise SystemExit("FATAL: mmdc not installed; the diagrams cannot be drawn")
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    side = {}
    for key, src in mermaid_blocks().items():
        with tempfile.TemporaryDirectory() as td:
            i = Path(td) / "d.mmd"
            i.write_text(src, encoding="utf-8")
            png = DIAGRAMS / f"E_{key}.png"
            r = subprocess.run([mmdc, "-i", str(i), "-o", str(png), "-w", "2400", "-b", "white",
                                "-q"], capture_output=True, text=True, timeout=180)
            if r.returncode != 0 or not png.is_file():
                raise SystemExit(f"FATAL: mmdc failed on diagram ({key}): {r.stderr[-800:]}")
        side[key] = {"mermaid_sha256": sha(src.encode()), "png": png.name,
                     "png_sha256": sha(png.read_bytes()), "drawn_by": f"mmdc ({mmdc})"}
    DIAGRAM_SIDECAR.write_text(json.dumps(side, indent=1, sort_keys=True) + "\n",
                               encoding="utf-8")
    print(f"drew {len(side)} diagrams under {DIAGRAMS.relative_to(REPO)}")
    return 0


def diagram_png(key: str) -> Path:
    side = json.loads(DIAGRAM_SIDECAR.read_text(encoding="utf-8"))
    src = mermaid_blocks()[key]
    ent = side.get(key)
    if ent is None or ent["mermaid_sha256"] != sha(src.encode()):
        raise DeckError(f"FATAL: diagram ({key}) is stale against E_architecture.md; run "
                        "--render-diagrams")
    png = DIAGRAMS / ent["png"]
    if sha(png.read_bytes()) != ent["png_sha256"]:
        raise DeckError(f"FATAL: {png.name} does not match its sidecar digest")
    return png


# ------------------------------------------------------------------------------ render

def display(text: str) -> str:
    """Backticks mark identifiers in the content file; the slide shows the identifier."""
    return text.replace("`", "")


def break_long(lines: list) -> list:
    """A single paragraph longer than a whole slide at the floor cannot be placed by the
    layout rule, which splits only between lines. Such a paragraph is broken at word boundaries
    into pieces that each fit; every word is kept, in order."""
    cap = FD.capacity(FD.FLOOR_PT) - 1
    per = max(20, int(FD.BODY_W * 72 / (0.5 * FD.FLOOR_PT)))
    out = []
    for ln in lines:
        if ln.blank or ln.pre or FD.wrapped_lines([ln], FD.FLOOR_PT) <= cap:
            out.append(ln)
            continue
        width = cap * max(10, per - 4 * ln.level) - 8
        cur = ""
        for word in ln.text.split(" "):
            if cur and len(cur) + 1 + len(word) > width:
                out.append(FD.Line(cur, level=ln.level))
                cur = word
            else:
                cur = f"{cur} {word}" if cur else word
        if cur:
            out.append(FD.Line(cur, level=ln.level))
    return out


def slide_plan(specs: list) -> list:
    """Each spec's lines, font and chunks, from the imported layout rule."""
    plan = []
    for spec in specs:
        lines = FD.parse_body(spec["body"])
        for ln in lines:
            ln.text = display(ln.text)
        lines = break_long(lines)
        if spec["layout"] == "cover" or spec["diagram"]:
            plan.append((spec, lines, None, [lines], False))
            continue
        pt, chunks, split = FD.layout(lines)
        plan.append((spec, lines, pt, chunks, split))
    return plan


def diagram_text_fits(lines: list, pt: int = FD.FLOOR_PT) -> bool:
    return FD.wrapped_lines(lines, pt) <= int(DIAGRAM_TEXT_H / (1.2 * pt / 72))


def _column_lines(lines: list, pt: int, width: float) -> int:
    """`FD.wrapped_lines`' estimate (0.5 em per character) for a column `width` inches wide
    rather than the full body width, which is the one number the side placement changes."""
    per = max(20, int(width * 72 / (0.5 * pt)))
    return sum(1 if ln.blank or ln.pre else max(1, -(-len(ln.text) // max(10, per - 4 * ln.level)))
               for ln in lines)


def diagram_placement(size: tuple, lines: list, pt: int = FD.FLOOR_PT) -> dict:
    """`{"mode", "scale", "img": (left, top, w, h), "text": (left, top, w, h)}` in inches:
    the placement that draws the image larger, side only when the text fits its column."""
    w, h = size
    stacked = min(FD.BODY_W / w, DIAGRAM_IMG_H / h)
    side_w = FD.BODY_W - DIAGRAM_SIDE_TEXT_W - DIAGRAM_GAP
    side = min(side_w / w, DIAGRAM_SIDE_H / h)
    side_cap = int(DIAGRAM_SIDE_H / (1.2 * pt / 72)) - FD.FIT_SLACK
    if side > stacked and _column_lines(lines, pt, DIAGRAM_SIDE_TEXT_W) <= side_cap:
        iw, ih = w * side, h * side
        return {"mode": "side", "scale": side,
                "img": (0.5 + (side_w - iw) / 2, 1.3, iw, ih),
                "text": (0.5 + side_w + DIAGRAM_GAP, 1.3, DIAGRAM_SIDE_TEXT_W, DIAGRAM_SIDE_H),
                "fits": True}
    iw, ih = w * stacked, h * stacked
    return {"mode": "stacked", "scale": stacked,
            "img": (0.5 + (FD.BODY_W - iw) / 2, 1.3, iw, ih),
            "text": (0.5, 1.3 + DIAGRAM_IMG_H + 0.1, FD.BODY_W, DIAGRAM_TEXT_H),
            "fits": diagram_text_fits(lines, pt)}


def build(specs: list, out: Path) -> dict:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    BLACK = RGBColor(0, 0, 0)
    GREY = RGBColor(0x55, 0x55, 0x55)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(FD.SLIDE_W), Inches(FD.SLIDE_H)
    report = {"slides": [], "splits": [], "overflow": []}

    def textbox(slide, top, height, lines, pt, bold_first=False, left=0.5, width=FD.BODY_W):
        bb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        bf = bb.text_frame
        bf.word_wrap = True
        for j, ln in enumerate(lines):
            p = bf.paragraphs[0] if j == 0 else bf.add_paragraph()
            if ln.blank:
                p.add_run().text = ""
                for r in p.runs:
                    r.font.size = Pt(max(8, pt // 2))
                continue
            p.level = 0 if ln.pre else ln.level
            FD.add_runs(p, ln.text, bold_all=bold_first and j == 0)
            for r in p.runs:
                r.font.size = Pt(FD.MONO_PT if ln.pre else pt)
                r.font.color.rgb = BLACK
                if ln.pre:
                    r.font.name = "Courier New"

    def title_and_footer(slide, title, sources):
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(FD.BODY_W), Inches(0.9))
        tb.text_frame.word_wrap = True
        FD.add_runs(tb.text_frame.paragraphs[0], display(title), bold_all=True)
        for r in tb.text_frame.paragraphs[0].runs:
            r.font.size, r.font.color.rgb = Pt(26), BLACK
        if sources:
            fb = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(FD.BODY_W),
                                          Inches(0.35))
            run = fb.text_frame.paragraphs[0].add_run()
            run.text = "source: " + ", ".join(sources)
            run.font.size, run.font.color.rgb = Pt(FOOTER_PT), GREY

    for spec, lines, pt, chunks, split in slide_plan(specs):
        if spec["layout"] == "cover":
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            body = [l for l in lines if not l.blank]
            tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(3.2))
            tb.text_frame.word_wrap = True
            FD.add_runs(tb.text_frame.paragraphs[0], body[0].text, bold_all=True)
            for r in tb.text_frame.paragraphs[0].runs:
                r.font.size, r.font.color.rgb = Pt(30), BLACK
            sb = slide.shapes.add_textbox(Inches(0.8), Inches(4.9), Inches(11.7), Inches(1.8))
            sb.text_frame.word_wrap = True
            for j, ln in enumerate(body[1:]):
                p = sb.text_frame.paragraphs[0] if j == 0 else sb.text_frame.add_paragraph()
                FD.add_runs(p, ln.text)
                for r in p.runs:
                    r.font.size, r.font.color.rgb = Pt(18), BLACK
            report["slides"].append({"n": spec["n"], "title": spec["title"], "pt": "cover",
                                     "chapter": spec["chapter"], "kind": spec["kind"]})
            continue
        if spec["diagram"]:
            png = diagram_png(spec["diagram"])
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            title_and_footer(slide, spec["title"], spec["sources"])
            from PIL import Image
            with Image.open(png) as im:
                place = diagram_placement(im.size, lines)
            il, it, iw, ih = place["img"]
            slide.shapes.add_picture(str(png), Inches(il), Inches(it), Inches(iw), Inches(ih))
            if not place["fits"]:
                report["overflow"].append(spec["n"])
            tl, tt, tw, th = place["text"]
            textbox(slide, tt, th, lines, FD.FLOOR_PT, left=tl, width=tw)
            report["slides"].append({"n": spec["n"], "title": spec["title"], "pt": FD.FLOOR_PT,
                                     "chapter": spec["chapter"], "kind": spec["kind"],
                                     "diagram": png.name, "placement": place["mode"],
                                     "scale": place["scale"]})
            continue
        if split:
            report["splits"].append({"slide": spec["n"], "parts": len(chunks), "pt": pt})
        for i, chunk in enumerate(chunks):
            if not FD.fits(chunk, pt):
                report["overflow"].append(spec["n"])
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            title_and_footer(slide, spec["title"] if i == 0 else f"{spec['title']} (cont.)",
                             spec["sources"])
            textbox(slide, 1.45, FD.BODY_H, chunk, pt)
            report["slides"].append({"n": spec["n"], "title": spec["title"], "pt": pt,
                                     "chapter": spec["chapter"], "kind": spec["kind"],
                                     "part": i + 1})
    buf = io.BytesIO()
    prs.save(buf)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(reproducible_zip(buf.getvalue()))
    return report


def reproducible_zip(data: bytes) -> bytes:
    """The same archive with every member's timestamp fixed; member order, names and bytes
    are unchanged."""
    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=ZIP_EPOCH)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            dst.writestr(zi, src.read(info.filename))
    return buf.getvalue()


def slide_count(specs: list) -> int:
    """Rendered slides for these specs, from the layout plan `build` follows."""
    return sum(len(chunks) for _, _, _, chunks, _ in slide_plan(specs))


def split_specs() -> tuple[list, list]:
    """`(brief, appendix)` specs. The content file is read twice: once to plan the appendix,
    whose rendered slide count the brief's closing `@stamp appendix` then states."""
    src = CONTENT.read_text(encoding="utf-8")

    def part(auth):
        brief = [s for s in auth if s["chapter"] != APPENDIX_CHAPTER]
        head = [s for s in auth if s["chapter"] == APPENDIX_CHAPTER]
        if not head or auth[len(brief):] != head:
            raise DeckError(f"FATAL: the {APPENDIX_CHAPTER} chapter must exist and be the "
                            "content file's last")
        return brief, head + generated(auth[-1]["n"] + 1)

    _, appendix = part(authored(src, appendix_slides=0))
    brief, appendix = part(authored(src, appendix_slides=slide_count(appendix)))
    return brief, appendix


def all_specs() -> list:
    brief, appendix = split_specs()
    return brief + appendix


def render(out: Path, appendix_out: Path) -> dict:
    """`{"brief": report, "appendix": report}`, one per output file."""
    brief, appendix = split_specs()
    return {"brief": build(brief, out), "appendix": build(appendix, appendix_out)}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--render-diagrams", action="store_true")
    a = ap.parse_args(argv)
    if a.render_diagrams:
        return render_diagrams()
    outputs = {"brief": DECK, "appendix": APPENDIX}
    if a.check:
        ok = True
        with tempfile.TemporaryDirectory() as td:
            tmp = {k: Path(td) / v.name for k, v in outputs.items()}
            rep = render(tmp["brief"], tmp["appendix"])
            for k, real in outputs.items():
                same = real.is_file() and real.read_bytes() == tmp[k].read_bytes()
                ok = ok and same and not rep[k]["overflow"]
                print(f"{len(rep[k]['slides'])} slides rendered; "
                      f"{'identical to' if same else 'DRIFT from'} {real.relative_to(REPO)}")
        return 0 if ok else 1
    rep = render(DECK, APPENDIX)
    src_per = Counter(s["chapter"] for s in all_specs())
    overflow = []
    for k, real in outputs.items():
        r = rep[k]
        per = Counter(s["chapter"] for s in r["slides"])
        print(f"wrote {real.relative_to(REPO)}: {len(r['slides'])} slides")
        for ch in dict.fromkeys(s["chapter"] for s in r["slides"]):
            print(f"  {str(ch):>9}: {src_per[ch]:>3} sections -> {per[ch]:>3} slides")
        print("  splits:", r["splits"] or "none")
        for s in r["slides"]:
            if s.get("diagram"):
                print(f"  diagram {s['diagram']}: {s['placement']}, scale {s['scale']:.5f} in/px")
        overflow += r["overflow"]
    if overflow:
        print("OVERFLOW on slides:", sorted(set(overflow)))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
