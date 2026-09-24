"""The brief deck is a rendered view of the pack: `cc_tasks/2026-09-22_brief_deck_assembly.md`.

What is asserted (decision 6), and why each is the right test:

* **Idempotence**: re-render and byte-compare, the guard the pack and every other generated view
  here uses. The archive's member timestamps are fixed by the renderer, so bytes are comparable.
* **The numeral gate** (decision 4) holds on the content file, and it REFUSES a typed number and
  an unfound quotation: a gate that is only ever seen passing has not been shown to gate.
* **Every `source:` file exists**, and the content file names the pack commit it was built from.
* **No slide overflows** the layout budget at the 14pt floor; oversize slides were split.
* **The appendix** is 49 indicator slides, plus one per rule group, plus one corpus summary.
* **Two outputs, nothing dropped** (`cc_tasks/2026-09-23_brief_deck_packaging.md` decision 1):
  the brief and the appendix together hold exactly the slides one file of the same sections
  would, the only sections the split added are the brief's closing slide and the appendix's
  cover, and the closing slide states the appendix's real slide count.
* **Chapter B packs to the floor** (decision 2): the skeleton §8 items are one section, and
  every part it splits into but the last is full.
* **Upright diagrams are drawn larger** (decision 3): (c) and (d) take the placement that draws
  them larger than the stacked box would, and their text fits it.
* **Chapter 0, the case** (`cc_tasks/2026-09-23_brief_narrative.md` decision 4): one slide per
  `## ` section plus the title slide, after the cover and before chapter A; every one carries
  notes equal to its section's paragraphs; the narrative passes the numeral and quotation gates,
  and the gate REFUSES a narrative number that is on no page its tag names.
* **The framework deck is unchanged** by the refactor that made its layout importable: rebuilt
  from its content file, every archive member equals the committed
  `docs/crosswalk/framework_deck_2026-09-02.pptx` (the zip header timestamps are the only bytes
  python-pptx does not reproduce).
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))

import build_brief_deck as D  # noqa: E402
import build_framework_deck as FD  # noqa: E402


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    d = tmp_path_factory.mktemp("deck")
    outs = {"brief": d / "brief.pptx", "appendix": d / "appendix.pptx"}
    return outs, D.render(outs["brief"], outs["appendix"])


@pytest.fixture(scope="module")
def specs():
    return D.all_specs()


def test_both_outputs_regenerate_byte_for_byte(rendered):
    outs, _ = rendered
    for key, real in (("brief", D.DECK), ("appendix", D.APPENDIX)):
        assert real.is_file(), f"{real} is not on disk"
        assert real.read_bytes() == outs[key].read_bytes(), \
            f"{real.name} drifted; run scripts/build_brief_deck.py"


def test_no_slide_overflows(rendered):
    _, reps = rendered
    for rep in reps.values():
        assert rep["overflow"] == [], f"slides over the layout budget: {rep['overflow']}"
        for s in rep["slides"]:
            assert s["pt"] in ("cover",) or s["pt"] >= FD.FLOOR_PT, s


def test_the_split_drops_nothing_and_adds_only_the_cover_and_the_closing_slide(rendered):
    _, reps = rendered
    brief, appendix = D.split_specs()
    assert len(reps["brief"]["slides"]) + len(reps["appendix"]["slides"]) \
        == D.slide_count(brief + appendix)
    assert [s["chapter"] for s in brief if s["chapter"] == "Close"] == ["Close"]
    assert brief[-1]["chapter"] == "Close"
    assert appendix[0]["layout"] == "cover" and appendix[0]["chapter"] == D.APPENDIX_CHAPTER
    assert all(s["chapter"] == D.APPENDIX_CHAPTER for s in appendix)
    assert all(s["chapter"] != D.APPENDIX_CHAPTER for s in brief)
    # The closing slide states the appendix file and the slide count that file really has.
    n = len(reps["appendix"]["slides"])
    assert f"{D.APPENDIX.relative_to(D.REPO)}: {n} slides" in brief[-1]["body"]


def test_the_skeleton_items_are_one_section_packed_to_the_floor():
    brief, _ = D.split_specs()
    items = [s for s in brief if s["chapter"] == "B" and "Decomposition with receipts" in s["body"]]
    assert len(items) == 1
    assert "11. **Red teaming" in items[0]["body"]
    (_, lines, pt, chunks, _), = D.slide_plan(items)
    assert pt == FD.FLOOR_PT or len(chunks) == 1
    for chunk, nxt in zip(chunks, chunks[1:]):
        assert not FD.fits(chunk + nxt[:1], pt), "a part was split before the floor filled it"


def test_upright_diagrams_take_the_placement_that_draws_them_larger(rendered):
    from PIL import Image
    _, reps = rendered
    placed = {s["diagram"]: s for s in reps["brief"]["slides"] if s.get("diagram")}
    for key in ("c", "d"):
        with Image.open(D.diagram_png(key)) as im:
            w, h = im.size
        assert h * 1.0 / w > 0.5, f"diagram ({key}) is not drawn upright: {w}x{h}"
        stacked = min(FD.BODY_W / w, D.DIAGRAM_IMG_H / h)
        s = placed[f"E_{key}.png"]
        assert s["placement"] == "side" and s["scale"] > stacked, s


def test_every_source_exists_and_every_number_stating_slide_names_one(specs):
    for s in specs:
        for src in s["sources"]:
            assert D.pack_path(src).is_file(), f"slide {s['n']}: {src}"
    assert all(s["sources"] for s in specs if s["kind"] != "authored")


def _content(body: str, n: int = 1, source: str = "C_provenance.md") -> str:
    return f"## Slide {n} — T\n\nsource: {source}\n{body}\n"


def test_the_numeral_gate_passes_a_ledger_number_and_refuses_a_typed_one():
    ledger = json.loads(D.LEDGER.read_text(encoding="utf-8"))
    held = ledger["C_provenance.md"][0]["value"]
    D.authored(_content(f"- {held} cells carry a locator"))
    with pytest.raises(SystemExit, match="numeral '4417'"):
        D.authored(_content("- 4417 cells carry a locator"))
    with pytest.raises(SystemExit, match="numeral"):
        D.authored(_content("- `4417` hidden in backticks"))


def test_a_quotation_must_be_found_in_the_named_source():
    D.authored(_content("> No admitted document matches `Title 13`, `CIPSEA`"))
    with pytest.raises(SystemExit, match="quote not found"):
        D.authored(_content("> 4417 documents were admitted and all are cited."))


def test_the_content_file_names_the_pack_it_was_built_from():
    first = D.CONTENT.read_text(encoding="utf-8").splitlines()[0]
    sha, _ = D.pack_commit()
    h = D.pack_header()
    assert f"brief pack at commit {sha}" in first, (first, sha)
    assert h["record_commit"] in first and h["cycle"] in first


def test_the_appendix_is_one_slide_per_sheet_per_rule_group_plus_the_corpus(specs):
    # Generated from the pack: chapter 0 (kind "case") is rendered from the narrative, not the pack.
    gen = [s for s in specs if s["kind"] in ("indicator", "rule", "corpus")]
    assert [s["kind"] for s in specs].index("indicator") == len(specs) - len(gen)
    sheets = list((D.PACK / "appendix").glob("indicator_*.md"))
    _, groups = D.rule_groups()
    per_rule = sum(len(r) for _, r in groups)
    rule_slides = per_rule if per_rule <= D.RULE_SLIDE_LIMIT else len(groups)
    assert len(sheets) == 49
    assert len([s for s in gen if s["kind"] == "indicator"]) == len(sheets)
    assert len([s for s in gen if s["kind"] == "rule"]) == rule_slides
    assert len(gen) == 49 + rule_slides + 1


def test_captured_output_stays_preformatted():
    lines, _ = D.directive("capture", "seldon_go", 0)
    parsed = FD.parse_body("\n".join(lines))
    assert all(l.pre for l in parsed[1:] if not l.blank), \
        [l.text for l in parsed[1:] if not l.pre and not l.blank][:3]


def test_the_diagrams_are_current_against_the_pack():
    for key in D.mermaid_blocks():
        assert D.diagram_png(key).is_file()


def test_the_framework_deck_is_unchanged_by_the_layout_refactor(tmp_path):
    out = tmp_path / "fd.pptx"
    r = subprocess.run([sys.executable, str(REPO / "scripts" / "build_framework_deck.py"),
                        "--out", str(out)], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    ref = zipfile.ZipFile(REPO / "docs" / "crosswalk" / "framework_deck_2026-09-02.pptx")
    got = zipfile.ZipFile(out)
    assert got.namelist() == ref.namelist()
    assert [n for n in ref.namelist() if got.read(n) != ref.read(n)] == []


def test_the_renderer_imports_the_layout_rule_rather_than_copying_it():
    src = (REPO / "scripts" / "build_brief_deck.py").read_text(encoding="utf-8")
    assert "FD.layout(" in src and "FD.parse_body(" in src
    assert not re.search(r"^def (parse_body|layout|fits|capacity|wrapped_lines)\b", src, re.M)


def test_chapter_0_is_one_slide_per_section_plus_the_title_after_the_cover(rendered):
    _, reps = rendered
    secs = D.parse_narrative(D.NARRATIVE.read_text(encoding="utf-8"))
    n_sections = len(secs) - 1
    slides = reps["brief"]["slides"]
    case = [s for s in slides if s["chapter"] == D.CASE_CHAPTER]
    assert len(case) == n_sections + 1
    chapters = [s["chapter"] for s in slides]
    first = chapters.index(D.CASE_CHAPTER)
    assert slides[0]["pt"] == "cover" and first == 1
    assert chapters[first + len(case)] == "A"


def test_every_chapter_0_slide_has_notes_equal_to_its_section_paragraphs(rendered):
    from pptx import Presentation
    outs, reps = rendered
    secs = D.parse_narrative(D.NARRATIVE.read_text(encoding="utf-8"))
    prs = Presentation(str(outs["brief"]))
    idx = [i for i, s in enumerate(reps["brief"]["slides"]) if s["chapter"] == D.CASE_CHAPTER]
    assert len(idx) == len(secs)
    for i, sec in zip(idx, secs):
        slide = prs.slides[i]
        assert slide.has_notes_slide, sec["title"]
        assert slide.notes_slide.notes_text_frame.text == "\n\n".join(sec["paras"]), sec["title"]
        assert sec["paras"], sec["title"]


def test_the_ask_projects_only_the_pending_ruling():
    ask = [s for s in D.case_specs() if s["title"] == D.ASK_TITLE]
    assert len(ask) == 1
    assert ask[0]["body"] == "- " + D.ASK_PENDING
    assert ask[0]["notes"].startswith("Draft, for the operator's ruling")


def test_chapter_tags_leave_the_slide_for_a_footer():
    for s in D.case_specs():
        assert not D.TAG_RE.search(s["body"]), s["title"]
    g = [s for s in D.case_specs() if s["title"] == "Start with our own product"][0]
    assert g["footer"].startswith("see chapter G")


def test_the_narrative_passes_the_numeral_and_quotation_gates():
    secs = D.parse_narrative(D.NARRATIVE.read_text(encoding="utf-8"))
    assert D.narrative_gate(secs) == []


def test_the_narrative_gate_refuses_a_number_and_a_quote_its_page_does_not_hold():
    held = json.loads(D.LEDGER.read_text(encoding="utf-8"))["G_census_dogfood.md"][0]["value"]
    ok = [{"title": "T", "bullets": [f"Score {held} [G]."], "paras": []}]
    assert D.narrative_gate(ok) == []
    bad = D.narrative_gate([{"title": "T", "bullets": ["4417 legs judged [G]."],
                             "paras": ["> 4417 documents were admitted and all are cited. [C]"]}])
    assert any("numeral '4417'" in b for b in bad), bad
    assert any("quote not found" in b for b in bad), bad
    # The number is held by page G but the sentence cites B: the tag decides the page.
    wrong = D.narrative_gate([{"title": "T", "bullets": [f"Score {held} [B]."], "paras": []}])
    assert wrong and "chapter(s) ['B']" in wrong[0], wrong


def test_the_narrative_names_its_record_cycle_and_task():
    first = D.NARRATIVE.read_text(encoding="utf-8").splitlines()[0]
    h = D.pack_header()
    assert first.startswith("<!--") and h["record_commit"] in first and h["cycle"] in first
    assert "cc_tasks/2026-09-23_brief_narrative.md at commit " in first


# ---------------------------------------------------------------- criterion labels (v3 decision 3)
# `cc_tasks/2026-09-24_brief_narrative_v3.md` decision 3: the narrative names a criterion by its
# letter and a label, and the numeral gate cannot see a wrong label or a count written as a
# criterion count USAFacts never made. Page B's criteria table is the one source of the labels.
# Three patterns, each designed against the sentences the narrative actually has:
#   * a range, "A through D" or "E to G", followed in its sentence by a list of exactly as many
#     items ("w, w, w and w"): the items, in order, are the range's labels;
#   * a lone letter A to G, not the first word of its sentence (there it is the article "A"), not
#     inside a range or a tag like "[B]", not an indicator code like "A10", followed by an
#     optional "the" and a phrase running to the next comma, colon, bracket, stop or "and": the
#     phrase is the letter's label, unless its first word is a connective or a verb (`_STOP`);
#   * a sentence naming "USAFacts" or "the guide", where the word before the first "criteria"
#     after that name is a number or a number word: it is four, and a four-item list after it
#     is A to D's labels.
_SENT = re.compile(r"(?<=[.\]])\s+(?=[A-Z0-9])")
_LETTER = re.compile(r"(?<![\w\[\]'’-])([A-G])(?![\w\]'’-])")
_RANGE = re.compile(r"(?<![\w\[\]'’-])([A-G]) (?:to|through) ([A-G])(?![\w\]'’-])")
_PHRASE = re.compile(r"\s+(?:the\s+)?(.+?)(?=\s*(?:[,;.:\[\]()]|\band\b|$))")
_STOP = {"to", "through", "and", "or", "are", "is", "was", "were", "has", "have", "had", "each",
         "all", "both", "carry", "carries", "fail", "fails", "pass", "passes", "which", "that"}
_ITEM = r"[\w'’-]+(?: [\w'’-]+){0,3}"
_LEAD = re.compile(r"^(?:(?:are|is|namely|the)\s+)+", re.I)
_NUMWORD = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7,
            "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12}
_SUBJECT = re.compile(r"USAFacts|\b[Tt]he guide\b")


def _page_b_labels() -> dict:
    text = (D.PACK / "B_usafacts_delta.md").read_text(encoding="utf-8")
    return {m.group(1): m.group(2).strip()
            for m in re.finditer(r"(?m)^\| ([A-G]) \| ([^|]+) \|", text)}


def _key(s: str) -> str:
    from kg.extraction.grounding import normalize
    return normalize(s).casefold()


def _a_list(text: str, n: int):
    m = re.search(rf"({_ITEM})" + rf", ({_ITEM})" * (n - 2) + rf",? and ({_ITEM})", text)
    return [_LEAD.sub("", x) for x in m.groups()] if m else None


def label_defects(narrative: str, labels: dict) -> list:
    """Every criterion label or USAFacts criterion count in `narrative` that page B contradicts."""
    bad = []
    for line in narrative.splitlines():
        if line.startswith("<!--") or not line.strip():
            continue
        for s in _SENT.split(line.lstrip("#>- ")):
            spans = []
            for m in _RANGE.finditer(s):
                spans.append(m.span())
                letters = [chr(c) for c in range(ord(m.group(1)), ord(m.group(2)) + 1)]
                got = _a_list(s[m.end():], len(letters)) if len(letters) > 1 else None
                if got and [_key(x) for x in got] != [_key(labels[c]) for c in letters]:
                    bad.append(f"{m.group(0)} labelled {got}, page B says "
                               f"{[labels[c] for c in letters]}: {s[:90]!r}")
            for m in _LETTER.finditer(s):
                if m.start() == 0 or any(a <= m.start() < b for a, b in spans):
                    continue
                pm = _PHRASE.match(s, m.end())
                phrase = pm.group(1).strip() if pm else ""
                if not phrase or phrase.split(" ")[0].casefold() in _STOP:
                    continue
                if _key(phrase) != _key(labels[m.group(1)]):
                    bad.append(f"{m.group(1)} labelled {phrase!r}, page B says "
                               f"{labels[m.group(1)]!r}: {s[:90]!r}")
            sm = _SUBJECT.search(s)
            cm = re.match(r".*?(\S+)\s+criteria\b", s[sm.end():]) if sm else None
            if cm:
                w = cm.group(1).casefold()
                n = int(w) if w.isdigit() else _NUMWORD.get(w)
                if n is not None and n != 4:
                    bad.append(f"USAFacts' criteria counted {n}, page B says four: {s[:90]!r}")
                got = _a_list(s[sm.end() + cm.end():], 4) if n == 4 else None
                if got and [_key(x) for x in got] != [_key(labels[c]) for c in "ABCD"]:
                    bad.append(f"USAFacts' four criteria listed {got}: {s[:90]!r}")
    return bad


def test_every_criterion_label_in_the_narrative_is_page_bs():
    labels = _page_b_labels()
    assert sorted(labels) == list("ABCDEFG"), labels
    assert label_defects(D.NARRATIVE.read_text(encoding="utf-8"), labels) == []


@pytest.mark.parametrize("current,v2", [
    ("USAFacts' guide gives agencies four criteria for AI-ready data, accessible, "
     "understandable, accurate and open, written",
     "USAFacts' guide gives agencies 7 criteria [G] for AI-ready data, written"),
    ("The guide calls itself a roadmap, and its four criteria, A through D, are accessible, "
     "understandable, accurate and open [B].",
     "Four of the criteria, A through D, describe the public surface a publisher controls: "
     "accessible, documented, licensed and cataloged."),
    ("Three criteria, E the TEVV loop, F release engineering and G the FSS-derived constructs, "
     "have", "Three criteria, E evaluation, F release and G governance, have"),
], ids=["v2_count_7", "v2_A_to_D_labels", "v2_E_to_G_labels"])
def test_the_label_gate_refuses_each_v2_error(current, v2):
    text = D.NARRATIVE.read_text(encoding="utf-8")
    assert text.count(current) == 1, current
    assert label_defects(text.replace(current, v2), _page_b_labels()), v2
